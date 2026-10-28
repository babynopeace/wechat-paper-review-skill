#!/usr/bin/env python3
"""Generate editable Chinese WeChat-paper-review DOCX and Markdown from article.json."""
from __future__ import annotations
import argparse
import json
import re
import os
from urllib.parse import quote
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


def load_article(path:Path)->dict:
    data=json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data,dict):
        raise ValueError("Expected an article JSON object")
    return data


def set_run_font(run,size=12,bold=False,latin="Times New Roman",cjk="宋体"):
    run.font.name=latin
    run.font.size=Pt(size)
    run.bold=bold
    rpr=run._element.get_or_add_rPr()
    rfonts=rpr.rFonts
    if rfonts is None:
        rfonts=OxmlElement('w:rFonts');rpr.insert(0,rfonts)
    rfonts.set(qn('w:eastAsia'),cjk)
    return run


def style_document(doc):
    sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.54);sec.bottom_margin=Cm(2.54)
    sec.left_margin=Cm(3.175);sec.right_margin=Cm(3.175)
    normal=doc.styles['Normal']
    normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.space_after=Pt(4)
    for name,size in [('Heading 1',14),('Heading 2',12)]:
        st=doc.styles[name]
        st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(11 if name=='Heading 1' else 7)
        st.paragraph_format.space_after=Pt(5)
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.line_spacing=1.15


def add_body(doc,text):
    para=doc.add_paragraph(style='Normal')
    para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.first_line_indent=Pt(24)
    para.paragraph_format.widow_control=True
    set_run_font(para.add_run(text))
    return para


def add_fig(doc,image,caption,num):
    if not image.exists():
        raise FileNotFoundError(f"Missing source figure: {image}")
    from PIL import Image
    with Image.open(image) as im:
        width,height=im.size
    maxw=14.25
    imageratio=width/max(1,height)
    # Avoid a single figure consuming multiple pages without warning.
    maxheight=18.0
    target_width=min(maxw,maxheight*imageratio)
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(7)
    p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(image),width=Cm(target_width))
    q=doc.add_paragraph()
    q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_after=Pt(9)
    q.paragraph_format.keep_together=True
    set_run_font(q.add_run(f"图{num}. {caption}"),size=8.5)


def each_block(section):
    for block in section.get('blocks',[]):
        yield block
    for sub in section.get('subsections',[]):
        for block in sub.get('blocks',[]):
            yield block


def export(data:dict, base:Path, outfile:Path, mdfile:Path|None=None):
    meta=data.get('metadata',{})
    d=Document();style_document(d)
    orig=meta.get('original_title','')
    p=d.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(10)
    set_run_font(p.add_run(orig),size=13,bold=True)
    for lab,key in [('作者','authors'),('发表期刊','journal'),('发表日期','publication_date')]:
        if meta.get(key):
            line=d.add_paragraph()
            line.paragraph_format.space_after=Pt(2)
            set_run_font(line.add_run(f"{lab}：{meta[key]}"),size=11)
    if meta.get('journal_metrics'):
        add_body(d,'期刊信息：'+str(meta['journal_metrics']))
    if meta.get('doi'):
        p=d.add_paragraph();set_run_font(p.add_run('DOI：'+str(meta['doi'])),size=10)
    for lead in data.get('lead',[]):
        add_body(d,lead)
    md=[f"# {meta.get('wechat_title') or orig}","",f"**原文标题：** {orig}","",f"**作者：** {meta.get('authors','')}","",f"**发表期刊：** {meta.get('journal','')}","",f"**发表日期：** {meta.get('publication_date','')}",""]
    if meta.get('doi'):
        md.extend([f"**DOI：** {meta['doi']}",""])
    for lead in data.get('lead',[]):
        md.extend([lead,""])
    figcount=0
    for i,sec in enumerate(data.get('sections',[]),1):
        p=d.add_paragraph(style='Heading 1')
        set_run_font(p.add_run(f"{i}. {sec['title']}"),size=14,bold=True)
        md.extend([f"## {i}. {sec['title']}",""])
        def insert(block):
            nonlocal figcount
            tp=block.get('type')
            if tp=='paragraph':
                add_body(d,block['text'])
                md.extend([block['text'],""])
            elif tp=='figure':
                figcount+=1
                source= (base / block['image']).resolve()
                # Never follow an arbitrary path from article JSON outside its workspace.
                if not source.is_relative_to(base.resolve()):
                    raise ValueError(f"Figure path escapes article directory: {source}")
                add_fig(d,source,block.get('caption',''),figcount)
                image_link = quote(Path(os.path.relpath(source, mdfile.parent.resolve() if mdfile else base.resolve())).as_posix(), safe='/')
                md.extend([f"![图{figcount}: {block.get('caption','')}]({image_link})", "",f"*图{figcount}. {block.get('caption','')}*", ""])
            else:
                raise ValueError(f"Unknown block type: {tp}")
        for block in sec.get('blocks',[]):
            insert(block)
        for j,sub in enumerate(sec.get('subsections',[]),1):
            p=d.add_paragraph(style='Heading 2')
            set_run_font(p.add_run(f"{i}.{j} {sub['title']}"),size=12,bold=True)
            md.extend([f"### {i}.{j} {sub['title']}",""])
            for block in sub.get('blocks',[]):
                insert(block)
    outfile.parent.mkdir(parents=True,exist_ok=True)
    d.save(outfile)
    if mdfile:
        mdfile.parent.mkdir(parents=True,exist_ok=True)
        mdfile.write_text('\n'.join(md).rstrip()+'\n',encoding='utf-8')
    return {"sections":len(data.get('sections',[])),"figures":figcount,"docx":str(outfile)}

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--article',required=True,type=Path)
    p.add_argument('--out',required=True,type=Path)
    p.add_argument('--markdown',type=Path)
    a=p.parse_args()
    r=export(load_article(a.article),a.article.parent,a.out,a.markdown)
    print(json.dumps(r,ensure_ascii=False))
